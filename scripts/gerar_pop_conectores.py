#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera o documento POP (Procedimento Operacional Padrão) da BLBW para
conectorização de fibra óptica com conector rápido SC/APC.

Uso:
    python3 scripts/gerar_pop_conectores.py

Saída:
    docs/POP-FTTH-001_Conectorizacao_Fibra_Optica_BLBW.docx
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Cores institucionais
AZUL_BLBW = RGBColor(0x1F, 0x3B, 0x73)      # azul escuro
VERDE_APC = RGBColor(0x2E, 0x7D, 0x32)      # verde (conector APC)
CINZA_CLARO = "D9E2F3"                       # fundo de cabecalho de tabela
AMARELO_ALERTA = "FFF2CC"                    # fundo de caixas de atencao
VERMELHO_PERIGO = "FDE9E9"                   # fundo de caixas de perigo


def set_cell_background(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "8EAADB")
        borders.append(el)
    tc_pr.append(borders)


def style_table(table, header_row=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_borders(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(10)
                    if header_row and i == 0:
                        run.font.bold = True
                        run.font.color.rgb = AZUL_BLBW
            if header_row and i == 0:
                set_cell_background(cell, CINZA_CLARO)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = AZUL_BLBW if level == 1 else VERDE_APC
    return h


def add_para(doc, text, bold=False, size=11, space_after=6, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.size = Pt(11)
        p.paragraph_format.space_after = Pt(3)


def add_steps(doc, steps, start=1):
    """Passos numerados com destaque no título do passo."""
    n = start
    for titulo, detalhe in steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"Passo {n} — {titulo}")
        run.font.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = AZUL_BLBW
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.75)
        p2.paragraph_format.space_after = Pt(8)
        run2 = p2.add_run(detalhe)
        run2.font.size = Pt(11)
        n += 1
    return n


def add_alert_box(doc, titulo, texto, cor_fundo=AMARELO_ALERTA):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_background(cell, cor_fundo)
    set_cell_borders(cell)
    p = cell.paragraphs[0]
    run = p.add_run(titulo + "\n")
    run.font.bold = True
    run.font.size = Pt(10.5)
    run2 = p.add_run(texto)
    run2.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BLBW — POP-FTTH-001 — Uso interno    |    Página ")
    run.font.size = Pt(8)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    run2.font.size = Pt(8)
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_end)


def build_document():
    doc = Document()

    # Estilo base
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    add_page_number_footer(section)

    # ------------------------------------------------------------------
    # Cabeçalho institucional (tabela de identificação do documento)
    # ------------------------------------------------------------------
    header = doc.add_table(rows=3, cols=4)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Linha 1: logo + título
    cell_logo = header.cell(0, 0)
    p = cell_logo.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BLBW")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = AZUL_BLBW

    cell_titulo = header.cell(0, 1).merge(header.cell(0, 3))
    p = cell_titulo.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PROCEDIMENTO OPERACIONAL PADRÃO (POP)\n")
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = AZUL_BLBW
    run2 = p.add_run("Conectorização de Fibra Óptica — Conector Rápido SC/APC")
    run2.font.size = Pt(12)
    run2.font.bold = True

    # Linha 2 e 3: metadados
    meta = [
        ("Código: POP-FTTH-001", "Versão: 1.0", "Data: 02/07/2026", "Páginas: —"),
        ("Elaborado por: Equipe Técnica", "Revisado por: Supervisão",
         "Aprovado por: Coordenação", "Área: Redes FTTH"),
    ]
    for r, valores in enumerate(meta, start=1):
        for c, valor in enumerate(valores):
            cell = header.cell(r, c)
            p = cell.paragraphs[0]
            run = p.add_run(valor)
            run.font.size = Pt(9)

    for row in header.rows:
        for cell in row.cells:
            set_cell_borders(cell)
    for cell in header.rows[0].cells:
        set_cell_background(cell, CINZA_CLARO)

    doc.add_paragraph()

    # ------------------------------------------------------------------
    # 1. Objetivo
    # ------------------------------------------------------------------
    add_heading(doc, "1. Objetivo")
    add_para(
        doc,
        "Este documento ensina, de forma simples e detalhada, como instalar um "
        "conector rápido SC/APC (aquele de corpo verde, montado em campo) na ponta "
        "de uma fibra óptica. O objetivo é que qualquer técnico da BLBW consiga "
        "fazer uma conexão de qualidade, com pouca perda de sinal, seguindo sempre "
        "o mesmo padrão."
    )

    # ------------------------------------------------------------------
    # 2. Onde este procedimento se aplica
    # ------------------------------------------------------------------
    add_heading(doc, "2. Onde este procedimento se aplica")
    add_bullets(doc, [
        "Instalação de clientes FTTH (fibra até a casa do cliente).",
        "Conectorização de cabo drop, cordão óptico ou fibra de tubo loose.",
        "Manutenção e troca de conectores danificados em campo ou no rack.",
    ])
    add_alert_box(
        doc,
        "IMPORTANTE — APC x UPC:",
        "O conector VERDE é APC (ponta polida em ângulo de 8°). O conector AZUL é "
        "UPC (ponta reta). Nunca conecte um APC em equipamento ou adaptador UPC: "
        "além de não funcionar direito, isso danifica as duas pontas. Na rede GPON "
        "da BLBW o padrão é SC/APC (verde).",
    )

    # ------------------------------------------------------------------
    # 3. Materiais e ferramentas
    # ------------------------------------------------------------------
    add_heading(doc, "3. Materiais e ferramentas necessárias")
    materiais = [
        ("Item", "Para que serve"),
        ("Conector rápido SC/APC (verde)", "Conector que será montado na ponta da fibra."),
        ("Decapador de fibra (stripper Miller ou de 3 furos)",
         "Retira as capas de proteção da fibra sem quebrá-la."),
        ("Clivador de precisão",
         "Faz o corte reto e perfeito na ponta da fibra. É a ferramenta mais "
         "importante do processo — sem ela a conexão não fica boa."),
        ("Álcool isopropílico (91% ou mais)", "Limpa a fibra antes do corte."),
        ("Lenços que não soltam fiapos (gaze ou kimwipes)", "Usados junto com o álcool na limpeza."),
        ("Alicate decapador de cabo drop e tesoura kevlar",
         "Abrem a capa externa do cabo, quando for cabo drop."),
        ("Power meter (medidor de potência óptica)",
         "Mede o sinal no final para confirmar que a conexão ficou boa."),
        ("Recipiente para descarte (garrafa ou caixinha fechada)",
         "Guarda os pedacinhos de fibra cortados, que são perigosos."),
        ("Óculos de proteção", "Protege os olhos contra estilhaços de fibra."),
    ]
    table = doc.add_table(rows=len(materiais), cols=2)
    for i, (col1, col2) in enumerate(materiais):
        table.cell(i, 0).paragraphs[0].add_run(col1)
        table.cell(i, 1).paragraphs[0].add_run(col2)
    style_table(table)
    doc.add_paragraph()

    # ------------------------------------------------------------------
    # 4. Segurança
    # ------------------------------------------------------------------
    add_heading(doc, "4. Segurança — leia antes de começar")
    add_alert_box(
        doc,
        "PERIGO — Estilhaços de fibra:",
        "Os pedacinhos de fibra cortados são finos como um fio de cabelo, "
        "transparentes e muito difíceis de enxergar. Eles furam a pele e, se "
        "entrarem no olho ou forem engolidos, podem causar ferimentos sérios. "
        "Por isso: use óculos de proteção, nunca sopre a fibra, não passe a mão "
        "na bancada e jogue TODO pedaço cortado no recipiente de descarte.",
        cor_fundo=VERMELHO_PERIGO,
    )
    add_alert_box(
        doc,
        "PERIGO — Luz invisível (laser):",
        "A luz que passa dentro da fibra é invisível e pode machucar o olho de "
        "forma permanente. NUNCA olhe diretamente para a ponta de uma fibra ou "
        "de um conector, mesmo que pareça 'apagado'.",
        cor_fundo=VERMELHO_PERIGO,
    )
    add_bullets(doc, [
        "Não coma nem beba enquanto trabalha com fibra.",
        "Lave as mãos ao terminar o serviço.",
        "Trabalhe em local iluminado e com superfície de cor escura, se possível "
        "(fica mais fácil enxergar os pedaços de fibra).",
    ])

    # ------------------------------------------------------------------
    # 5. Procedimento passo a passo
    # ------------------------------------------------------------------
    add_heading(doc, "5. Procedimento passo a passo")

    add_heading(doc, "5.1. Etapa 1 — Preparar o conector", level=2)
    n = add_steps(doc, [
        ("Confira o modelo do conector",
         "Verifique se o conector é SC/APC (corpo verde) e leia o manual ou o "
         "gabarito que vem na embalagem. Cada fabricante indica o tamanho exato "
         "de fibra que deve ficar exposta (geralmente entre 10 e 13 mm)."),
        ("Desmonte o conector",
         "Separe as partes: o corpo verde (frente), a capa traseira (tail cap) e "
         "a bota de proteção. Deixe tudo à mão, em cima de uma superfície limpa."),
        ("Passe as peças traseiras na fibra ANTES de tudo",
         "Enfie primeiro a bota e depois a capa traseira no cabo/fibra que será "
         "conectorizado. ATENÇÃO: este é o erro mais comum de todos! Se você "
         "clivar a fibra antes de passar essas peças, não conseguirá mais "
         "colocá-las e terá que refazer todo o trabalho."),
    ])

    add_heading(doc, "5.2. Etapa 2 — Preparar a fibra", level=2)
    n = add_steps(doc, [
        ("Abra o cabo (se for cabo drop ou tubo loose)",
         "Com o alicate decapador, retire a capa externa do cabo deixando uns "
         "5 a 6 cm de folga para trabalhar. Se a fibra vier dentro de tubo com "
         "gel, limpe todo o gel com lenço e álcool isopropílico. Corte o "
         "sustentador (arame ou kevlar) com a tesoura apropriada."),
        ("Decape a fibra",
         "Com o stripper, retire o revestimento colorido da fibra (a casquinha "
         "de 250 µm e, se houver, o buffer de 900 µm), deixando de 3 a 4 cm de "
         "fibra 'nua' (vidro transparente). Dica: retire aos poucos, cerca de "
         "1 cm por vez, segurando a fibra reta. Se puxar tudo de uma vez, a "
         "fibra pode quebrar."),
        ("Limpe a fibra nua",
         "Molhe o lenço com álcool isopropílico e passe na fibra nua sempre no "
         "mesmo sentido (da capa para a ponta), 2 ou 3 vezes. Quando a fibra "
         "está limpa, ela faz um barulhinho de 'rangido' ao passar o lenço. "
         "Depois de limpa, NÃO toque mais na fibra com os dedos e não deixe "
         "ela encostar em nada."),
    ], start=n)

    add_heading(doc, "5.3. Etapa 3 — Clivar a fibra (o corte perfeito)", level=2)
    add_para(
        doc,
        "A clivagem é o momento mais importante de todo o procedimento. Um corte "
        "torto ou lascado é a principal causa de conexão ruim. Capriche aqui.",
    )
    n = add_steps(doc, [
        ("Posicione a fibra no clivador",
         "Coloque a fibra no clivador usando a régua/gabarito do próprio "
         "clivador. A medida é contada a partir do fim do revestimento (onde "
         "termina a parte colorida) até a lâmina. Use a medida indicada pelo "
         "fabricante do conector — na maioria dos conectores rápidos SC/APC "
         "fica entre 10 e 12 mm de fibra nua."),
        ("Faça o corte",
         "Feche a tampa do clivador e acione a lâmina com movimento firme e "
         "único. Não force e não repita o movimento. O clivador faz um corte "
         "perfeitamente reto (90°), que é o que garante a qualidade da conexão."),
        ("Cuide da ponta cortada",
         "Depois do corte, a ponta da fibra está pronta e NÃO PODE encostar em "
         "nada — nem no dedo, nem na mesa, nem na roupa. Não sopre e não limpe "
         "a ponta. Pegue o pedacinho de fibra que sobrou no clivador e jogue "
         "imediatamente no recipiente de descarte."),
    ], start=n)

    add_heading(doc, "5.4. Etapa 4 — Montar o conector", level=2)
    n = add_steps(doc, [
        ("Abra a trava do conector",
         "Destrave o mecanismo interno do conector (cunha, tampa ou clipe, "
         "dependendo do modelo) para permitir a entrada da fibra."),
        ("Insira a fibra até sentir o encosto",
         "Segure a fibra bem perto do conector e vá inserindo pelo furo "
         "traseiro, devagar, sempre reta e SEM girar. Empurre até sentir que a "
         "ponta encostou em algo firme lá dentro (é o pedacinho de fibra que já "
         "vem de fábrica dentro do conector, chamado stub). Continue empurrando "
         "levemente até a fibra formar uma pequena curva (barriga) do lado de "
         "fora. Essa curva é o sinal de que a ponta está bem encostada, com "
         "pressão de contato."),
        ("Trave a fibra",
         "Mantendo essa pequena curva, feche a trava do conector para prender a "
         "fibra. Depois de travar, a curva pode diminuir um pouco, mas a fibra "
         "não pode ficar solta. Se a fibra escorregar para fora, abra a trava e "
         "repita a inserção."),
        ("Finalize a montagem",
         "Rosqueie ou encaixe a capa traseira no corpo do conector e deslize a "
         "bota até cobrir a emenda. Se o conector tiver braçadeira para prender "
         "o cabo drop ou a fibra 900 µm, prenda conforme o modelo. Pronto: o "
         "conector está montado."),
    ], start=n)

    add_heading(doc, "5.5. Etapa 5 — Testar a conexão", level=2)
    n = add_steps(doc, [
        ("Meça o sinal",
         "Conecte o conector no adaptador/equipamento e meça o sinal com o "
         "power meter, ou verifique o nível de RX na ONU do cliente. Na rede "
         "GPON, o sinal recebido deve ficar entre -8 dBm e -27 dBm. O ideal é "
         "que o conector rápido adicione menos de 0,5 dB de perda (quanto "
         "menor, melhor)."),
        ("Se o sinal estiver ruim, refaça",
         "Sinal fraco ou ausente quase sempre significa clivagem ruim ou fibra "
         "que não encostou direito no stub. Solução: desmonte o conector, corte "
         "1 a 2 cm da fibra e refaça tudo a partir da limpeza (Etapa 2, Passo "
         "6). Não tente 'ajeitar' um conector mal montado — refazer é mais "
         "rápido e mais confiável."),
    ], start=n)

    # ------------------------------------------------------------------
    # 6. Erros comuns
    # ------------------------------------------------------------------
    add_heading(doc, "6. Erros mais comuns e como evitar")
    erros = [
        ("Erro", "Consequência", "Como evitar"),
        ("Esquecer de passar a bota e a capa traseira antes de clivar",
         "Precisa cortar a fibra e refazer tudo",
         "Sempre passe as peças na fibra logo no início (Passo 3)."),
        ("Clivar com a medida errada",
         "A fibra não alcança o stub — sem sinal",
         "Use o gabarito do conector e confira a medida antes de cortar."),
        ("Tocar ou sujar a ponta depois de clivada",
         "Perda alta de sinal",
         "Depois do corte, a ponta não encosta em nada. Se encostar, clive de novo."),
        ("Girar a fibra ao inserir no conector",
         "A ponta pode lascar",
         "Insira sempre reto, devagar, sem girar."),
        ("Travar sem a fibra estar encostada no stub (sem a curvinha)",
         "Fica um espaço de ar dentro do conector — perda e reflexão altas",
         "Só trave quando ver a pequena curva na fibra."),
        ("Conectar APC (verde) em UPC (azul)",
         "Danifica as duas pontas e degrada o sinal",
         "Verde com verde, azul com azul. Na dúvida, confira a cor."),
    ]
    table = doc.add_table(rows=len(erros), cols=3)
    for i, (c1, c2, c3) in enumerate(erros):
        table.cell(i, 0).paragraphs[0].add_run(c1)
        table.cell(i, 1).paragraphs[0].add_run(c2)
        table.cell(i, 2).paragraphs[0].add_run(c3)
    style_table(table)
    doc.add_paragraph()

    # ------------------------------------------------------------------
    # 7. Descarte
    # ------------------------------------------------------------------
    add_heading(doc, "7. Descarte dos resíduos")
    add_bullets(doc, [
        "Junte TODOS os pedaços de fibra cortados no recipiente de descarte "
        "fechado (garrafa PET ou caixinha rígida com tampa).",
        "Nunca jogue pedaços de fibra soltos no lixo comum, no chão ou no bolso.",
        "Ao final do serviço, confira a bancada e o chão à procura de pedaços "
        "esquecidos.",
        "Descarte o recipiente cheio conforme a orientação da BLBW para resíduos "
        "perfurocortantes.",
    ])

    # ------------------------------------------------------------------
    # 8. Resumo rápido
    # ------------------------------------------------------------------
    add_heading(doc, "8. Resumo rápido (cola de bolso)")
    resumo = [
        ("Ordem", "Ação"),
        ("1", "Passar bota + capa traseira na fibra."),
        ("2", "Decapar 3 a 4 cm de fibra nua."),
        ("3", "Limpar com álcool isopropílico até 'ranger'."),
        ("4", "Clivar na medida do gabarito (10 a 12 mm)."),
        ("5", "Inserir reto, sem girar, até formar a curvinha."),
        ("6", "Travar, fechar capa traseira e bota."),
        ("7", "Medir o sinal. Ruim? Cortar e refazer."),
    ]
    table = doc.add_table(rows=len(resumo), cols=2)
    for i, (c1, c2) in enumerate(resumo):
        table.cell(i, 0).paragraphs[0].add_run(c1)
        table.cell(i, 1).paragraphs[0].add_run(c2)
    style_table(table)
    doc.add_paragraph()

    # ------------------------------------------------------------------
    # 9. Histórico de revisões
    # ------------------------------------------------------------------
    add_heading(doc, "9. Histórico de revisões")
    revisoes = [
        ("Versão", "Data", "Descrição", "Responsável"),
        ("1.0", "02/07/2026", "Emissão inicial do documento.", "Equipe Técnica BLBW"),
    ]
    table = doc.add_table(rows=len(revisoes), cols=4)
    for i, valores in enumerate(revisoes):
        for c, valor in enumerate(valores):
            table.cell(i, c).paragraphs[0].add_run(valor)
    style_table(table)

    return doc


def main():
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "POP-FTTH-001_Conectorizacao_Fibra_Optica_BLBW.docx")
    doc = build_document()
    doc.save(out_path)
    print(f"Documento gerado: {out_path}")


if __name__ == "__main__":
    main()
