#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gerador DOCX e PDF - Manual TIA Portal BLBW."""

import os
import sys
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from fpdf import FPDF

sys.path.insert(0, os.path.dirname(__file__))
from manual_content import MANUAL_META, INDEX, CHAPTERS
from manual_content_extended import (
    EXTENDED_INDEX,
    EXTENDED_CHAPTERS,
    SECTION_ADDITIONS,
)


def build_manual_data():
    """Mescla conteúdo base com extensões e adições por seção."""
    import copy

    chapters = copy.deepcopy(CHAPTERS)

    for chapter in chapters:
        for section in chapter["sections"]:
            title = section["title"]
            if title in SECTION_ADDITIONS:
                section["instructions"].extend(SECTION_ADDITIONS[title])

    chapters.extend(EXTENDED_CHAPTERS)
    index = list(INDEX) + list(EXTENDED_INDEX)
    return index, chapters


INDEX_FULL, CHAPTERS_FULL = build_manual_data()

# BLBW Corporate Colors
BLBW_PRIMARY = RGBColor(0x00, 0x66, 0xB3)      # #0066B3
BLBW_DARK = RGBColor(0x00, 0x33, 0x66)           # #003366
BLBW_ACCENT = RGBColor(0x00, 0xA0, 0xE3)         # #00A0E3
BLBW_TEXT = RGBColor(0x33, 0x33, 0x33)

PDF_PRIMARY = (0, 102, 179)
PDF_DARK = (0, 51, 102)
PDF_ACCENT = (0, 160, 227)
PDF_TEXT = (51, 51, 51)

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "docs")


def set_cell_shading(cell, color_hex):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def setup_docx_styles(doc):
    styles = doc.styles
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLBW_TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for level, size, color in [(1, 18, BLBW_DARK), (2, 14, BLBW_PRIMARY), (3, 12, BLBW_DARK)]:
        h = styles[f"Heading {level}"]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = color
        h.paragraph_format.space_before = Pt(12 if level > 1 else 0)
        h.paragraph_format.space_after = Pt(6)

    if "Code Block" not in [s.name for s in styles]:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(9)
        code_style.paragraph_format.left_indent = Cm(0.5)
        code_style.paragraph_format.space_before = Pt(4)
        code_style.paragraph_format.space_after = Pt(4)


def add_docx_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = f"{MANUAL_META['company']}  |  {MANUAL_META['classification']}"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = BLBW_PRIMARY
        run.font.italic = True

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run(f"Versão {MANUAL_META['version']}  —  ")
    run.font.size = Pt(8)
    run.font.color.rgb = BLBW_TEXT
    run = fp.add_run("Página ")
    run.font.size = Pt(8)
    run.font.color.rgb = BLBW_TEXT
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    run2 = fp.add_run(" de ")
    run2.font.size = Pt(8)
    run2.font.color.rgb = BLBW_TEXT
    fld2 = OxmlElement("w:fldSimple")
    fld2.set(qn("w:instr"), "NUMPAGES")
    fp._p.append(fld2)


def add_cover_page(doc):
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BLBW")
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = BLBW_PRIMARY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Be Life, Be Water")
    run.font.size = Pt(14)
    run.font.color.rgb = BLBW_ACCENT

    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(MANUAL_META["title"])
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = BLBW_DARK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(MANUAL_META["subtitle"])
    run.font.size = Pt(13)
    run.font.color.rgb = BLBW_TEXT

    doc.add_paragraph()
    doc.add_paragraph()

    info = [
        ("Versão:", MANUAL_META["version"]),
        ("Público-Alvo:", MANUAL_META["audience"]),
        ("Plataforma:", MANUAL_META["platform"]),
        ("Data:", datetime.now().strftime("%d/%m/%Y")),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = 1
    for i, (label, value) in enumerate(info):
        row = table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = value
        for cell in row.cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(11)
                    run.font.color.rgb = BLBW_TEXT

    doc.add_page_break()


def add_index(doc):
    doc.add_heading("ÍNDICE GERAL", level=1)
    for chapter, sections in INDEX_FULL:
        p = doc.add_paragraph()
        run = p.add_run(chapter)
        run.bold = True
        run.font.color.rgb = BLBW_PRIMARY
        run.font.size = Pt(12)
        for sec in sections:
            sp = doc.add_paragraph(sec, style="List Bullet")
            sp.paragraph_format.left_indent = Cm(1)
    doc.add_page_break()


def add_instruction_docx(doc, instr):
    p = doc.add_heading(instr["name"], level=3)

    fields = [
        ("Nome e Objetivo", instr["objective"]),
        ("Quando usar", instr["when_use"]),
        ("Sintaxe e Parâmetros", instr["syntax"]),
    ]
    for label, text in fields:
        p = doc.add_paragraph()
        run = p.add_run(f"{label}: ")
        run.bold = True
        run.font.color.rgb = BLBW_DARK
        p.add_run(text)

    p = doc.add_paragraph()
    run = p.add_run("Exemplo Prático de Aplicação:")
    run.bold = True
    run.font.color.rgb = BLBW_DARK

    p = doc.add_paragraph()
    run = p.add_run("LAD (Ladder):")
    run.bold = True
    run.italic = True
    for line in instr["lad"].split("\n"):
        doc.add_paragraph(line, style="Code Block")

    p = doc.add_paragraph()
    run = p.add_run("STL (Statement List):")
    run.bold = True
    run.italic = True
    for line in instr["stl"].split("\n"):
        doc.add_paragraph(line, style="Code Block")

    doc.add_paragraph()


def generate_docx(output_path):
    doc = Document()
    setup_docx_styles(doc)
    add_docx_header_footer(doc)
    add_cover_page(doc)
    add_index(doc)

    for chapter in CHAPTERS_FULL:
        doc.add_heading(chapter["title"], level=1)
        for section in chapter["sections"]:
            doc.add_heading(section["title"], level=2)
            for instr in section["instructions"]:
                add_instruction_docx(doc, instr)

    doc.save(output_path)
    print(f"DOCX gerado: {output_path}")


FONT_DIR = "/usr/share/fonts/truetype/dejavu"


class BLBW_PDF(FPDF):
    def setup_fonts(self):
        self.add_font("DejaVu", "", f"{FONT_DIR}/DejaVuSans.ttf")
        self.add_font("DejaVu", "B", f"{FONT_DIR}/DejaVuSans-Bold.ttf")
        self.add_font("DejaVu", "I", f"{FONT_DIR}/DejaVuSans.ttf")
        self.add_font("DejaVu", "BI", f"{FONT_DIR}/DejaVuSans-Bold.ttf")
        self.add_font("DejaVuMono", "", f"{FONT_DIR}/DejaVuSansMono.ttf")

    def header(self):
        self.set_font("DejaVu", "I", 8)
        self.set_text_color(*PDF_PRIMARY)
        self.cell(0, 8, f"{MANUAL_META['company']}  |  {MANUAL_META['classification']}", align="R")
        self.ln(4)
        self.set_draw_color(*PDF_PRIMARY)
        self.line(10, self.get_y(), 200, self.get_y())
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("DejaVu", "I", 8)
        self.set_text_color(*PDF_TEXT)
        self.cell(0, 10, f"Versão {MANUAL_META['version']}  -  Página {self.page_no()}/{{nb}}", align="C")

    def chapter_title(self, title):
        self.set_font("DejaVu", "B", 16)
        self.set_text_color(*PDF_DARK)
        self.cell(0, 12, title, new_x="LMARGIN", new_y="NEXT")
        self.ln(2)

    def section_title(self, title):
        self.set_font("DejaVu", "B", 13)
        self.set_text_color(*PDF_PRIMARY)
        self.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
        self.ln(1)

    def instruction_title(self, title):
        self.set_font("DejaVu", "B", 11)
        self.set_text_color(*PDF_DARK)
        self.multi_cell(0, 6, title)
        self.ln(1)

    def body_text(self, label, text):
        self.set_font("DejaVu", "B", 10)
        self.set_text_color(*PDF_DARK)
        self.write(5, f"{label}: ")
        self.set_font("DejaVu", "", 10)
        self.set_text_color(*PDF_TEXT)
        self.multi_cell(0, 5, text)
        self.ln(2)

    def code_block(self, text):
        self.set_font("DejaVuMono", "", 8)
        self.set_fill_color(245, 245, 245)
        self.set_text_color(30, 30, 30)
        for line in text.split("\n"):
            self.cell(0, 4, line, new_x="LMARGIN", new_y="NEXT", fill=True)
        self.ln(3)


def generate_pdf(output_path):
    pdf = BLBW_PDF()
    pdf.setup_fonts()
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    # Cover
    pdf.ln(30)
    pdf.set_font("DejaVu", "B", 36)
    pdf.set_text_color(*PDF_PRIMARY)
    pdf.cell(0, 15, "BLBW", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 14)
    pdf.set_text_color(*PDF_ACCENT)
    pdf.cell(0, 8, "Be Life, Be Water", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(15)
    pdf.set_font("DejaVu", "B", 18)
    pdf.set_text_color(*PDF_DARK)
    pdf.multi_cell(0, 8, MANUAL_META["title"], align="C")
    pdf.ln(5)
    pdf.set_font("DejaVu", "", 12)
    pdf.set_text_color(*PDF_TEXT)
    pdf.multi_cell(0, 6, MANUAL_META["subtitle"], align="C")
    pdf.ln(15)
    for label, value in [
        ("Versão:", MANUAL_META["version"]),
        ("Público-Alvo:", MANUAL_META["audience"]),
        ("Plataforma:", MANUAL_META["platform"]),
        ("Data:", datetime.now().strftime("%d/%m/%Y")),
    ]:
        pdf.set_font("DejaVu", "B", 10)
        pdf.cell(40, 7, label, align="R")
        pdf.set_font("DejaVu", "", 10)
        pdf.cell(0, 7, value, new_x="LMARGIN", new_y="NEXT")

    # Index
    pdf.add_page()
    pdf.chapter_title("ÍNDICE GERAL")
    for chapter, sections in INDEX_FULL:
        pdf.set_font("DejaVu", "B", 11)
        pdf.set_text_color(*PDF_PRIMARY)
        pdf.cell(0, 7, chapter, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "", 10)
        pdf.set_text_color(*PDF_TEXT)
        for sec in sections:
            pdf.cell(0, 5, f"    - {sec}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    # Chapters
    for chapter in CHAPTERS_FULL:
        pdf.add_page()
        pdf.chapter_title(chapter["title"])
        for section in chapter["sections"]:
            pdf.section_title(section["title"])
            for instr in section["instructions"]:
                if pdf.get_y() > 240:
                    pdf.add_page()
                pdf.instruction_title(instr["name"])
                pdf.body_text("Nome e Objetivo", instr["objective"])
                pdf.body_text("Quando usar", instr["when_use"])
                pdf.body_text("Sintaxe e Parâmetros", instr["syntax"])
                pdf.set_font("DejaVu", "B", 10)
                pdf.set_text_color(*PDF_DARK)
                pdf.cell(0, 6, "Exemplo Prático de Aplicação:", new_x="LMARGIN", new_y="NEXT")
                pdf.set_font("DejaVu", "BI", 9)
                pdf.cell(0, 5, "LAD (Ladder):", new_x="LMARGIN", new_y="NEXT")
                pdf.code_block(instr["lad"])
                pdf.set_font("DejaVu", "BI", 9)
                pdf.cell(0, 5, "STL (Statement List):", new_x="LMARGIN", new_y="NEXT")
                pdf.code_block(instr["stl"])
                pdf.ln(3)

    pdf.output(output_path)
    print(f"PDF gerado: {output_path}")


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    docx_path = os.path.join(OUTPUT_DIR, "Manual_Tecnico_TIA_Portal_BLBW.docx")
    pdf_path = os.path.join(OUTPUT_DIR, "Manual_Tecnico_TIA_Portal_BLBW.pdf")

    generate_docx(docx_path)
    generate_pdf(pdf_path)

    docx_size = os.path.getsize(docx_path)
    pdf_size = os.path.getsize(pdf_path)
    print(f"\nArquivos gerados em: {os.path.abspath(OUTPUT_DIR)}")
    print(f"  DOCX: {docx_size:,} bytes")
    print(f"  PDF:  {pdf_size:,} bytes")


if __name__ == "__main__":
    main()
