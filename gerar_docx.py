"""Gera a instrução de trabalho BLBW de fusão de fibra óptica (FSM-60S) em .docx."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x1F, 0x4E, 0x79)
CINZA = RGBColor(0x59, 0x59, 0x59)


def sombrear(celula, cor_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), cor_hex)
    celula._tc.get_or_add_tcPr().append(shd)


doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

estilo = doc.styles["Normal"]
estilo.font.name = "Calibri"
estilo.font.size = Pt(11)

# ---------------- Cabeçalho institucional ----------------
header = doc.sections[0].header
tab_h = header.add_table(rows=1, cols=3, width=Cm(16))
tab_h.alignment = WD_TABLE_ALIGNMENT.CENTER
c1, c2, c3 = tab_h.rows[0].cells
c1.text = "BLBW"
c1.paragraphs[0].runs[0].font.size = Pt(20)
c1.paragraphs[0].runs[0].font.bold = True
c1.paragraphs[0].runs[0].font.color.rgb = AZUL
c2.text = "INSTRUÇÃO DE TRABALHO\nFusão de Fibra Óptica – Fujikura FSM-60S"
for p in c2.paragraphs:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.bold = True
        r.font.size = Pt(10)
c3.text = "IT-FIB-001\nRev. 00"
for p in c3.paragraphs:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(9)

# ---------------- Título ----------------
titulo = doc.add_paragraph()
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = titulo.add_run("INSTRUÇÃO DE TRABALHO\nFUSÃO DE FIBRA ÓPTICA COM MÁQUINA FUJIKURA FSM-60S")
r.font.size = Pt(16)
r.font.bold = True
r.font.color.rgb = AZUL

# ---------------- Controle do documento ----------------
doc.add_paragraph()
tab = doc.add_table(rows=4, cols=4)
tab.style = "Table Grid"
dados = [
    ("Código:", "IT-FIB-001", "Revisão:", "00"),
    ("Elaborado por:", "", "Data:", "02/07/2026"),
    ("Aprovado por:", "", "Data:", ""),
    ("Área:", "Redes Ópticas / Infraestrutura", "Páginas:", "1 de 3"),
]
for linha, valores in zip(tab.rows, dados):
    for cel, valor in zip(linha.cells, valores):
        cel.text = valor
        if valor.endswith(":"):
            cel.paragraphs[0].runs[0].font.bold = True
            sombrear(cel, "D9E2F3")


def secao(num, texto):
    p = doc.add_paragraph()
    p.space_before = Pt(14)
    r = p.add_run(f"{num}. {texto}")
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = AZUL
    return p


def passo(num, titulo_passo, detalhes):
    p = doc.add_paragraph()
    r = p.add_run(f"Passo {num} – {titulo_passo}")
    r.font.bold = True
    for d in detalhes:
        doc.add_paragraph(d, style="List Bullet")


# ---------------- 1. Objetivo ----------------
secao(1, "OBJETIVO")
doc.add_paragraph(
    "Padronizar o procedimento de emenda por fusão de fibras ópticas utilizando a máquina de "
    "fusão Fujikura FSM-60S (alinhamento pelo núcleo), garantindo emendas com perda típica "
    "igual ou inferior a 0,02 dB e segurança do executante."
)

# ---------------- 2. Materiais e equipamentos ----------------
secao(2, "MATERIAIS E EQUIPAMENTOS NECESSÁRIOS")
for item in [
    "Máquina de fusão Fujikura FSM-60S (bateria carregada ou fonte ADC-13);",
    "Clivador de precisão;",
    "Decapador de fibra óptica (alicate stripper);",
    "Tubetes protetores de emenda (60 mm ou 40 mm);",
    "Álcool isopropílico (pureza ≥ 99%) e lenços/gaze sem fiapos;",
    "Power meter óptico (para conferência do enlace após a emenda);",
    "Óculos de proteção e descarte adequado para resíduos de fibra.",
]:
    doc.add_paragraph(item, style="List Bullet")

# ---------------- 3. Segurança ----------------
secao(3, "CUIDADOS DE SEGURANÇA")
for item in [
    "Nunca olhe diretamente para a extremidade de uma fibra que possa estar energizada com laser;",
    "Descarte os fragmentos de fibra clivada em recipiente apropriado — são estilhaços de vidro;",
    "Não toque nos eletrodos da máquina durante ou logo após o arco de fusão;",
    "Trabalhe em local limpo e protegido de vento e poeira sempre que possível.",
]:
    doc.add_paragraph(item, style="List Bullet")

# ---------------- 4. Procedimento ----------------
secao(4, "PROCEDIMENTO PASSO A PASSO")

passo(1, "Preparação inicial", [
    "Ligue a máquina e confirme o modo de fusão adequado (SM AUTO para fibra monomodo padrão G.652; AUTO se o tipo de fibra for desconhecido);",
    "Insira o TUBETE PROTETOR em uma das fibras ANTES de iniciar a preparação — após a fusão não é mais possível colocá-lo.",
])
passo(2, "Decapagem da fibra", [
    "Com o decapador, remova a capa e o acrilato deixando cerca de 3 a 4 cm de fibra nua;",
    "Remova em pequenos trechos para não quebrar a fibra.",
])
passo(3, "Limpeza", [
    "Limpe a fibra nua com gaze/lenço sem fiapos umedecido em álcool isopropílico;",
    "A fibra limpa produz um som característico ('canto') ao ser limpa;",
    "Não toque na fibra após a limpeza.",
])
passo(4, "Clivagem", [
    "Posicione a fibra no clivador respeitando o comprimento de clivagem (16 mm para capa de 250 µm com o grampo padrão);",
    "Acione a lâmina com movimento firme e único;",
    "NÃO limpe e NÃO toque na ponta clivada.",
])
passo(5, "Posicionamento na máquina", [
    "Abra a tampa corta-vento e as presilhas;",
    "Apoie a fibra na canaleta em V (V-groove), com a ponta entre a borda da canaleta e o centro dos eletrodos, sem encostar em nada;",
    "Feche a presilha. Repita os passos 2 a 5 para a segunda fibra, no lado oposto.",
])
passo(6, "Fusão", [
    "Feche a tampa corta-vento e pressione a tecla SET;",
    "A máquina executa automaticamente: arco de limpeza, verificação do ângulo de clive, alinhamento dos núcleos, arco de fusão e estimativa de perda;",
    "Resultado esperado: perda estimada ≤ 0,02 dB para fibra monomodo;",
    "Em caso de erro (ângulo de clive ruim, bolha, poeira), pressione RESET, clive novamente a fibra e repita.",
])
passo(7, "Retirada e proteção da emenda", [
    "Abra a tampa e retire a fibra com cuidado — a máquina aplica um teste de tração automático (~2 N);",
    "Deslize o tubete até centralizá-lo exatamente sobre o ponto da emenda;",
    "Mantenha a fibra levemente tensionada para não dobrá-la.",
])
passo(8, "Termocontração (forno)", [
    "Posicione o conjunto no forno na parte traseira da máquina, com a emenda centralizada;",
    "Pressione HEAT (ou aguarde o acionamento automático). O ciclo dura cerca de 30 s;",
    "Aguarde o sinal sonoro, retire e deixe esfriar na bandeja de resfriamento (J-Plate) antes de acomodar a fibra.",
])
passo(9, "Verificação final", [
    "Inspecione visualmente o tubete contraído (sem bolhas ou fibra exposta);",
    "Meça a atenuação do enlace com o power meter e registre o resultado;",
    "Acomode a emenda na bandeja/caixa de emenda respeitando o raio mínimo de curvatura.",
])

# ---------------- 5. Erros comuns ----------------
secao(5, "ERROS COMUNS E CORREÇÕES")
tab2 = doc.add_table(rows=5, cols=2)
tab2.style = "Table Grid"
erros = [
    ("Problema", "Ação corretiva"),
    ("Erro de ângulo de clive", "Clivar novamente; verificar lâmina do clivador"),
    ("Bolha na emenda", "Refazer limpeza e clivagem; verificar sujeira no V-groove"),
    ("Perda estimada alta", "Limpar V-groove e lentes; refazer preparação da fibra"),
    ("Arco fraco / fusões falhando", "Executar calibração de arco no menu; verificar desgaste dos eletrodos (vida útil ~2500 arcos)"),
]
for linha, (a, b) in zip(tab2.rows, erros):
    linha.cells[0].text = a
    linha.cells[1].text = b
for cel in tab2.rows[0].cells:
    cel.paragraphs[0].runs[0].font.bold = True
    sombrear(cel, "1F4E79")
    cel.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# ---------------- 6. Resumo ----------------
secao(6, "RESUMO RÁPIDO")
p = doc.add_paragraph()
r = p.add_run("TUBETE → DECAPA → LIMPA → CLIVA → POSICIONA → SET → CENTRALIZA TUBETE → HEAT → CONFERE")
r.font.bold = True
r.font.color.rgb = AZUL

# ---------------- Rodapé ----------------
footer = doc.sections[0].footer
pf = footer.paragraphs[0]
pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = pf.add_run("BLBW – Documento interno | IT-FIB-001 Rev. 00 | Reprodução proibida sem autorização")
rf.font.size = Pt(8)
rf.font.color.rgb = CINZA

doc.save("/workspace/IT-FIB-001_Fusao_Fibra_Optica_FSM-60S_BLBW.docx")
print("Documento gerado com sucesso.")
